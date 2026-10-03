import os
import uuid
import shutil
import zipfile
import subprocess
from pathlib import Path

from flask import Flask, request, send_file, render_template, jsonify, after_this_request
from werkzeug.utils import secure_filename
from pypdf import PdfReader, PdfWriter
import fitz
from PIL import Image

BASE = Path(__file__).parent
UPLOAD = BASE / "uploads"
OUT = BASE / "outputs"
UPLOAD.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024

TOOLS = [
    ("merge-pdf", "Merge PDF", "Combine multiple PDF files into one document", "📎"),
    ("split-pdf", "Split PDF", "Extract every page as a separate PDF", "✂️"),
    ("compress-pdf", "Compress PDF", "Reduce PDF file size", "🗜️"),
    ("pdf-to-word", "PDF to Word", "Convert PDF text into an editable DOCX", "W"),
    ("word-to-pdf", "Word to PDF", "Convert DOC and DOCX files to PDF", "W"),
    ("pdf-to-excel", "PDF to Excel", "Extract PDF text into an XLSX workbook", "X"),
    ("excel-to-pdf", "Excel to PDF", "Convert XLS and XLSX files to PDF", "X"),
    ("pdf-to-ppt", "PDF to PowerPoint", "Create a PowerPoint slide for each PDF page", "P"),
    ("ppt-to-pdf", "PowerPoint to PDF", "Convert PPT and PPTX files to PDF", "P"),
    ("jpg-to-pdf", "JPG to PDF", "Turn one or more images into a PDF", "🖼️"),
    ("pdf-to-jpg", "PDF to JPG", "Export PDF pages as JPG images", "🖼️"),
    ("rotate-pdf", "Rotate PDF", "Rotate every page by 90, 180 or 270 degrees", "↻"),
    ("watermark-pdf", "Watermark PDF", "Add text to every PDF page", "T"),
    ("protect-pdf", "Protect PDF", "Password-protect a PDF", "🔒"),
    ("unlock-pdf", "Unlock PDF", "Remove a known PDF password", "🔓"),
]


def save_upload(f, job):
    name = secure_filename(f.filename or "file")
    p = UPLOAD / f"{job}_{name}"
    f.save(p)
    return p


def cleanup_files(paths):
    for p in paths:
        try:
            Path(p).unlink(missing_ok=True)
        except Exception:
            pass


def send_result(path, filename):
    path = Path(path)

    @after_this_request
    def _cleanup(response):
        cleanup_files([path])
        return response

    return send_file(path, as_attachment=True, download_name=filename)


def zip_files(paths, out):
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            z.write(p, p.name)
    return out


def find_libreoffice():
    candidates = [
        shutil.which("libreoffice"),
        shutil.which("soffice"),
        "/usr/bin/libreoffice",
        "/usr/bin/soffice",
        r"C:\\Program Files\\LibreOffice\\program\\soffice.exe",
        r"C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe",
    ]
    for item in candidates:
        if item and Path(item).exists():
            return item
    return None


def office_to_pdf(src, out_dir):
    exe = find_libreoffice()
    if not exe:
        raise RuntimeError("Office conversion engine is unavailable on this server.")
    out_dir = Path(out_dir)
    out_dir.mkdir(exist_ok=True)
    profile = out_dir / f"lo_profile_{uuid.uuid4().hex}"
    profile.mkdir(exist_ok=True)
    try:
        cmd = [
            exe,
            "--headless",
            "--convert-to", "pdf",
            "--outdir", str(out_dir),
            "-env:UserInstallation=file:///" + str(profile).replace("\\", "/"),
            str(src),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
        expected = out_dir / (Path(src).stem + ".pdf")
        if result.returncode != 0 or not expected.exists():
            detail = (result.stderr or result.stdout or "conversion failed").strip()
            raise RuntimeError(f"Office conversion failed: {detail[:500]}")
        return expected
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def pdf_to_docx(src, out):
    from docx import Document
    docx = Document()
    pdf = fitz.open(str(src))
    try:
        for i, page in enumerate(pdf):
            text = page.get_text("text").strip()
            if i:
                docx.add_page_break()
            if text:
                for line in text.splitlines():
                    if line.strip():
                        docx.add_paragraph(line.strip())
    finally:
        pdf.close()
    docx.save(out)
    return out


def pdf_to_ppt(src, out):
    from pptx import Presentation
    from pptx.util import Inches

    pdf = fitz.open(str(src))
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    tmp_images = []
    try:
        for i, page in enumerate(pdf):
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            image = OUT / f"{uuid.uuid4().hex}_slide_{i}.jpg"
            pix.save(str(image))
            tmp_images.append(image)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
        prs.save(out)
    finally:
        pdf.close()
        cleanup_files(tmp_images)
    return out


@app.get("/")
def home():
    return render_template("index.html", tools=TOOLS)


@app.get("/health")
def health():
    return jsonify(status="ok")


@app.get("/tool/<slug>")
def tool(slug):
    item = next((x for x in TOOLS if x[0] == slug), None)
    if not item:
        return "Not found", 404
    return render_template("tool.html", tool=item)


@app.errorhandler(413)
def too_large(_):
    return jsonify(error="File too large. Maximum upload size is 100 MB."), 413


@app.post("/api/<slug>")
def api(slug):
    job = uuid.uuid4().hex[:10]
    files = request.files.getlist("files")
    if not files or all(not f.filename for f in files):
        return jsonify(error="Please select file(s)."), 400

    paths = [save_upload(f, job) for f in files if f.filename]
    generated = []
    try:
        if slug == "merge-pdf":
            writer = PdfWriter()
            for p in paths:
                writer.append(str(p))
            out = OUT / f"{job}_merged.pdf"
            writer.write(out)
            generated.append(out)
            return send_result(out, "merged.pdf")

        if slug == "split-pdf":
            reader = PdfReader(str(paths[0]))
            outs = []
            for i, page in enumerate(reader.pages):
                p = OUT / f"{job}_page_{i + 1}.pdf"
                writer = PdfWriter()
                writer.add_page(page)
                writer.write(p)
                outs.append(p)
            out = OUT / f"{job}_split.zip"
            zip_files(outs, out)
            generated.extend(outs + [out])
            @after_this_request
            def _cleanup_split(response):
                cleanup_files(generated)
                return response
            return send_file(out, as_attachment=True, download_name="split-pages.zip")

        if slug == "compress-pdf":
            doc = fitz.open(str(paths[0]))
            out = OUT / f"{job}_compressed.pdf"
            doc.save(str(out), garbage=4, deflate=True, clean=True)
            doc.close()
            generated.append(out)
            return send_result(out, "compressed.pdf")

        if slug == "pdf-to-word":
            out = OUT / f"{job}.docx"
            pdf_to_docx(paths[0], out)
            generated.append(out)
            return send_result(out, "converted.docx")

        if slug == "word-to-pdf":
            out = office_to_pdf(paths[0], OUT)
            generated.append(out)
            return send_result(out, "converted.pdf")

        if slug == "pdf-to-excel":
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "PDF Text"
            doc = fitz.open(str(paths[0]))
            try:
                row = 1
                for page in doc:
                    for line in page.get_text("text").splitlines():
                        if line.strip():
                            ws.cell(row=row, column=1, value=line.strip())
                            row += 1
            finally:
                doc.close()
            out = OUT / f"{job}.xlsx"
            wb.save(out)
            generated.append(out)
            return send_result(out, "converted.xlsx")

        if slug == "excel-to-pdf":
            out = office_to_pdf(paths[0], OUT)
            generated.append(out)
            return send_result(out, "converted.pdf")

        if slug == "pdf-to-ppt":
            out = OUT / f"{job}.pptx"
            pdf_to_ppt(paths[0], out)
            generated.append(out)
            return send_result(out, "converted.pptx")

        if slug == "ppt-to-pdf":
            out = office_to_pdf(paths[0], OUT)
            generated.append(out)
            return send_result(out, "converted.pdf")

        if slug == "jpg-to-pdf":
            images = []
            try:
                for p in paths:
                    images.append(Image.open(p).convert("RGB"))
                out = OUT / f"{job}.pdf"
                images[0].save(str(out), save_all=True, append_images=images[1:])
            finally:
                for im in images:
                    im.close()
            generated.append(out)
            return send_result(out, "images.pdf")

        if slug == "pdf-to-jpg":
            doc = fitz.open(str(paths[0]))
            outs = []
            try:
                for i, page in enumerate(doc):
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                    p = OUT / f"{job}_{i + 1}.jpg"
                    pix.save(str(p))
                    outs.append(p)
            finally:
                doc.close()
            out = OUT / f"{job}_jpg.zip"
            zip_files(outs, out)
            generated.extend(outs + [out])
            @after_this_request
            def _cleanup_jpg(response):
                cleanup_files(generated)
                return response
            return send_file(out, as_attachment=True, download_name="pdf-pages.zip")

        if slug == "rotate-pdf":
            deg = int(request.form.get("degrees", "90"))
            reader = PdfReader(str(paths[0]))
            writer = PdfWriter()
            for page in reader.pages:
                page.rotate(deg)
                writer.add_page(page)
            out = OUT / f"{job}_rotated.pdf"
            writer.write(out)
            generated.append(out)
            return send_result(out, "rotated.pdf")

        if slug == "watermark-pdf":
            text = request.form.get("text", "CONFIDENTIAL").strip() or "CONFIDENTIAL"
            doc = fitz.open(str(paths[0]))
            out = OUT / f"{job}_watermarked.pdf"
            try:
                for page in doc:
                    page.insert_text((72, 72), text, fontsize=22, fill=(0.6, 0.6, 0.6), overlay=True)
                doc.save(str(out))
            finally:
                doc.close()
            generated.append(out)
            return send_result(out, "watermarked.pdf")

        if slug == "protect-pdf":
            password = request.form.get("password", "")
            if not password:
                return jsonify(error="Enter a password."), 400
            reader = PdfReader(str(paths[0]))
            writer = PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            writer.encrypt(password)
            out = OUT / f"{job}_protected.pdf"
            writer.write(out)
            generated.append(out)
            return send_result(out, "protected.pdf")

        if slug == "unlock-pdf":
            password = request.form.get("password", "")
            reader = PdfReader(str(paths[0]))
            if reader.is_encrypted and not reader.decrypt(password):
                return jsonify(error="Incorrect password."), 400
            writer = PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            out = OUT / f"{job}_unlocked.pdf"
            writer.write(out)
            generated.append(out)
            return send_result(out, "unlocked.pdf")

        return jsonify(error="Tool not implemented."), 400

    except subprocess.TimeoutExpired:
        return jsonify(error="Conversion took too long. Try a smaller file."), 500
    except Exception as e:
        return jsonify(error=str(e)), 500
    finally:
        cleanup_files(paths)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
